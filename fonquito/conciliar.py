"""Concilia el formato institucional (.xlsm/.xlsx, Anexo 1 ingresos / Anexo 2 gastos) contra las facturas validadas.
Solo LEE el formato (no ejecuta macros ni lo modifica)."""
import re
from datetime import date, datetime

from openpyxl import load_workbook

TOL = 0.005


def _num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def _fecha(v):
    if isinstance(v, (datetime, date)):
        return v.date() if isinstance(v, datetime) else v
    try:
        return datetime.strptime(str(v).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def leer_anexo(ruta, tipo):
    """Filas ya digitadas del Anexo 1 (venta) o Anexo 2 (gasto): lista de dicts por nombre de columna."""
    wb = load_workbook(ruta, data_only=True, keep_vba=False)
    hoja = next((n for n in wb.sheetnames if n.upper().replace(" ", "") == ("ANEXO1" if tipo == "venta" else "ANEXO2")), None)
    if hoja is None:
        raise ValueError(f"No se encontró la hoja {'ANEXO 1' if tipo == 'venta' else 'ANEXO 2'} en {ruta}")
    ws = wb[hoja]
    cab = None
    filas = []
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400), values_only=True):
        if cab is None:
            if any(str(c).strip().upper() == "# FACTURA" for c in fila if c):
                cab = [str(c).strip().upper() if c else "" for c in fila]
            continue
        d = {cab[i]: v for i, v in enumerate(fila) if i < len(cab) and cab[i]}
        if re.fullmatch(r"\d{3}-\d{3}-\d{9}", str(d.get("# FACTURA") or "").strip()):
            filas.append(d)
    return filas


def conciliar(regs, filas, tipo):
    """Devuelve lista de (nivel, N°, factura, mensaje). nivel: ERROR | FALTA | SOBRA | INFO."""
    idx = {str(f["# FACTURA"]).strip(): f for f in filas}
    out, vistos = [], set()
    col_prov = "CLIENTE" if tipo == "venta" else "PROVEEDOR"
    col_ruc = "RUC" if tipo == "venta" else "RUC/DNI"
    for r in regs:
        num = r.get("numero")
        if r.get("duplicado_de"):
            continue
        f = idx.get(num)
        if f is None:
            out.append(("FALTA", r["n"], num, f"No está en el formato: {r.get('receptor') if tipo == 'venta' else r.get('proveedor')}, "
                                              f"{r['fecha_date']:%d/%m/%Y}, total ${r['total']:.2f}"))
            continue
        vistos.add(num)
        for campo, val in (("SUBTOTAL", r["subtotal"]), ("IVA", r["iva"] or 0.0), ("TOTAL", r["total"])):
            dig = _num(f.get(campo))
            if val is None:
                continue                      # la factura no se pudo leer: ya consta como advertencia
            if dig is None or abs(dig - val) > TOL:
                out.append(("ERROR", r["n"], num, f"{campo}: formato {dig} ≠ factura {val:.2f}"))
        if _fecha(f.get("FECHA")) != r["fecha_date"]:
            out.append(("ERROR", r["n"], num, f"FECHA: formato {f.get('FECHA')} ≠ factura {r['fecha_date']:%d/%m/%Y}"))
        ruc_f = str(f.get(col_ruc) or "").strip()
        ruc_r = (r.get("id_receptor") if tipo == "venta" else r.get("ruc_clave")) or ""
        if ruc_f and ruc_r and ruc_f != ruc_r:
            out.append(("ERROR", r["n"], num, f"{col_ruc}: formato {ruc_f} ≠ factura {ruc_r}"))
        if tipo == "gasto" and str(f.get("TIPO GASTO") or "").strip() and r.get("tipo_gasto") != str(f["TIPO GASTO"]).strip():
            out.append(("INFO", r["n"], num, f"TIPO GASTO: formato '{f['TIPO GASTO']}' / herramienta '{r.get('tipo_gasto')}' (se respeta el del formato)"))
    for num, f in idx.items():
        if num not in vistos and not any(r.get("numero") == num for r in regs):
            out.append(("SOBRA", None, num, f"Está en el formato pero no hay factura en el lote: {f.get(col_prov)}, total {f.get('TOTAL')}"))
    tot_f = round(sum(_num(f.get("TOTAL")) or 0 for f in filas), 2)
    tot_v = round(sum(r["total"] or 0 for r in regs if r["estado"] != "INVÁLIDA ❌"), 2)
    out.append(("INFO", None, "", f"Total del formato ${tot_f:.2f} ({len(filas)} filas) frente a lo validado ${tot_v:.2f} "
                                  f"({sum(1 for r in regs if r['estado'] != 'INVÁLIDA ❌')} facturas); diferencia ${tot_f - tot_v:.2f}"))
    return out
