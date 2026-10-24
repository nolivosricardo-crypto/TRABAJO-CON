import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from fonquito import sri
from fonquito.excel import construir
import shutil
import subprocess

from fonquito.extraer import dividir_paginas, extraer_campos, leer_md
from fonquito.renombrar import extraer_paginas, nombre_nuevo, renombrar
from fonquito.validar import Parametros, clasificar_gasto, huecos_numeracion, resumen, validar_lote

RUC = "1790000000001"          # RUC ficticio del beneficiario
EMISOR = "1791111111001"       # RUC ficticio del proveedor
DESDE, HASTA = date(2026, 8, 1), date(2026, 8, 31)


def clave(dia, secuencial, ruc=EMISOR):
    return sri.generar_clave(date(2026, 8, dia), "01", ruc, "2", "001001", f"{secuencial:09d}")


def factura_md(archivo, cl, numero, fecha, sub, iva, total, receptor="EMPRENDEDORA PRUEBA", idr=RUC,
               prov="PROVEEDOR PRUEBA S.A.", desc="INSUMOS DE COCINA", ruc=EMISOR):
    return f"""
## File: Recibidas/{archivo}

| | | R.U.C.: | {ruc} |
FACTURA
| No. | {numero} |
NÚMERO DE AUTORIZACIÓN
{cl}
| {prov} | FECHA Y HORA DE |
{fecha} 10:00:00
AUTORIZACIÓN:
Razón Social / Nombres y Apellidos: {receptor}
| Identificación | {idr} |
| Fecha | {fecha} | Placa / Matrícula: | Guía |
Principal Auxiliar Subsidio
1.00 {desc} 10.00 0.00 0.00 0.00 {sub:.2f}
SUBTOTAL 15% {sub:.2f}
SUBTOTAL SIN IMPUESTOS {sub:.2f}
IVA 15% {iva:.2f}
VALOR TOTAL {total:.2f}
"""


def md_lote(*bloques):
    return "Content from the zip file `X\\Recibidas.zip`:\n" + "".join(bloques) + "\n## File: __MACOSX/Recibidas/._.DS_Store\n\n"


class TestSRI(unittest.TestCase):
    def test_modulo11_roundtrip(self):
        c = clave(5, 1)
        info = sri.descomponer(c)
        self.assertTrue(info["modulo11_ok"])
        self.assertEqual(info["ruc"], EMISOR)
        self.assertEqual(info["numero"], "001-001-000000001")

    def test_digito_incorrecto(self):
        c = clave(5, 1)
        mal = c[:-1] + str((int(c[-1]) + 1) % 10)
        self.assertFalse(sri.descomponer(mal)["modulo11_ok"])

    def test_formato_invalido(self):
        with self.assertRaises(ValueError):
            sri.descomponer("123")

    def test_residuos_especiales(self):
        # residuo 0 -> 0 y residuo 1 -> 1 (nunca 11 ni 10)
        for base in ("0" * 48, "1" * 48):
            dv, _, _ = sri.modulo11(base)
            self.assertIn(dv, range(0, 10))


class TestPipeline(unittest.TestCase):
    def setUp(self):
        self.p = Parametros("gasto", RUC, "EMPRENDEDORA PRUEBA", DESDE, HASTA)
        c1, c2, c3, c4 = clave(3, 1), clave(10, 2), clave(12, 3), clave(20, 4)
        malo = c4[:-1] + str((int(c4[-1]) + 1) % 10)
        fuera = sri.generar_clave(date(2026, 9, 2), "01", EMISOR, "2", "001001", f"{5:09d}")
        self.md = md_lote(
            factura_md("a.pdf", c1, "001-001-000000001", "03/08/2026", 10.00, 1.50, 11.50),
            factura_md("b.pdf", c2, "001-001-000000002", "10/08/2026", 20.00, 3.00, 23.00, desc="SNACK TAKIS"),
            factura_md("b_copia.pdf", c2, "001-001-000000002", "10/08/2026", 20.00, 3.00, 23.00, desc="SNACK TAKIS"),
            factura_md("c.pdf", c3, "001-001-000000003", "12/08/2026", 10.00, 1.50, 12.00),            # monto mal
            factura_md("d.pdf", malo, "001-001-000000004", "20/08/2026", 5.00, 0.75, 5.75),            # dígito mal
            factura_md("e.pdf", fuera, "001-001-000000005", "02/09/2026", 5.00, 0.75, 5.75),          # fuera de período
            factura_md("f.pdf", clave(25, 6), "001-001-000000006", "25/08/2026", 5.00, 0.75, 5.75, idr="1790999999001"),  # receptor
        )
        self.tmp = tempfile.TemporaryDirectory()
        self.md_path = Path(self.tmp.name) / "r.md"
        self.md_path.write_text(self.md, encoding="utf-8")

    def tearDown(self):
        self.tmp.cleanup()

    def regs(self):
        docs = leer_md(self.md_path)
        self.assertNotIn(".DS_Store", " ".join(docs))
        return validar_lote([extraer_campos(t, n) for n, t in docs.items()], self.p)

    def test_extraccion(self):
        c = extraer_campos(leer_md(self.md_path)["a.pdf"], "a.pdf")
        self.assertEqual((c["numero"], c["fecha"], c["subtotal"], c["iva"], c["total"]),
                         ("001-001-000000001", "03/08/2026", 10.0, 1.5, 11.5))
        self.assertEqual(c["proveedor"], "PROVEEDOR PRUEBA S.A.")
        self.assertEqual(c["id_receptor"], RUC)

    def test_estados(self):
        regs = {r["archivo"]: r for r in self.regs()}
        self.assertEqual(regs["a.pdf"]["estado"], "VÁLIDA ✅")
        self.assertEqual(regs["b.pdf"]["estado"], "REVISAR ⚠️")           # palabra 'takis'
        self.assertEqual(regs["b_copia.pdf"]["estado"], "INVÁLIDA ❌")     # duplicado
        self.assertIsNotNone(regs["b_copia.pdf"]["duplicado_de"])
        for malo in ("c.pdf", "d.pdf", "e.pdf", "f.pdf"):
            self.assertEqual(regs[malo]["estado"], "INVÁLIDA ❌", malo)
        self.assertIn("MONTO INCONSISTENTE", " ".join(regs["c.pdf"]["observaciones"]))
        self.assertIn("Dígito verificador", " ".join(regs["d.pdf"]["observaciones"]))
        self.assertIn("FUERA DE PERÍODO", " ".join(regs["e.pdf"]["observaciones"]))
        self.assertIn("RECEPTOR INCORRECTO", " ".join(regs["f.pdf"]["observaciones"]))

    def test_proveedor_dos_lineas_formato_pdf(self):
        cl = clave(3, 1)
        texto = ("R.U.C.: 1791111111001\nFACTURA\nNo. 001-001-000000001\nNÚMERO DE AUTORIZACIÓN\n" + cl +
                 "\nEMPRESA DE PRUEBA PARA LA INDUSTRIA FECHA Y HORA DE\nDE ALIMENTOS S.A. AUTORIZACIÓN: 03/08/2026 13:20:07\n"
                 "EMPRESA DE PRUEBA PARA LA INDUSTRIA\nDE ALIMENTOS S.A. AMBIENTE: PRODUCCIÓN\n")
        self.assertEqual(extraer_campos(texto, "x.pdf")["proveedor"], "EMPRESA DE PRUEBA PARA LA INDUSTRIA DE ALIMENTOS S.A.")

    def test_revision_manual(self):
        docs = leer_md(self.md_path)
        regs = {r["archivo"]: r for r in validar_lote([extraer_campos(t, n) for n, t in docs.items()], self.p,
                                                      {"a.pdf": "Compra inusual", "c.pdf": "no debe pisar inválida"})}
        self.assertEqual(regs["a.pdf"]["estado"], "REVISAR ⚠️")
        self.assertIn("Compra inusual", regs["a.pdf"]["observaciones"])
        self.assertEqual(regs["c.pdf"]["estado"], "INVÁLIDA ❌")

    def test_totales_excluyen_invalidas(self):
        r = resumen(self.regs())
        self.assertEqual((r["procesados"], r["validas"], r["revisar"], r["invalidas"]), (7, 1, 1, 5))
        self.assertEqual(r["total"], round(11.50 + 23.00, 2))

    def test_orden_cronologico_y_numeracion(self):
        regs = self.regs()
        fechas = [r["fecha_date"] for r in regs]
        self.assertEqual(fechas, sorted(fechas))
        self.assertEqual([r["n"] for r in regs], list(range(1, len(regs) + 1)))

    def test_excel_estructura_y_formulas(self):
        regs = self.regs()
        wb = construir(regs, self.p, "PRUEBA", "Gastos_Ago2026_PRUEBA.xlsx", "Agosto 2026")
        ruta = Path(self.tmp.name) / "x.xlsx"
        wb.save(ruta)
        wb2 = load_workbook(ruta)
        self.assertEqual(wb2.sheetnames, ["Registro de Gastos", "Validación Módulo 11", "Renombrado"])
        ws = wb2["Registro de Gastos"]
        formulas = [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=SUMIF")]
        self.assertEqual(len(formulas), 9)  # subtotal/IVA/total x (total, excluidas, revisar)
        encabezados = [c.value for c in ws[9]]
        self.assertEqual(encabezados[:4], ["No.", "FECHA", "RUC/DNI", "# FACTURA"])
        self.assertIn("TIPO GASTO", encabezados)

    def test_renombrado(self):
        regs = self.regs()
        nuevos = [nombre_nuevo(r) for r in regs]
        self.assertEqual(len(set(nuevos)), len(nuevos))  # sin colisiones
        self.assertTrue(any(n.startswith("0") and "DUPLICADO" in n for n in nuevos))
        origen = Path(self.tmp.name) / "origen.zip"
        with zipfile.ZipFile(origen, "w") as z:
            for r in regs:
                if r["archivo"] != "f.pdf":
                    z.writestr(f"Recibidas/{r['archivo']}", b"%PDF-1.4 prueba")
            z.writestr("__MACOSX/Recibidas/._a.pdf", b"x")
        destino = Path(self.tmp.name) / "salida.zip"
        ok, falta = renombrar(regs, origen, destino)
        self.assertEqual(falta, ["f.pdf"])
        with zipfile.ZipFile(destino) as z:
            self.assertEqual(sorted(z.namelist()), sorted(n for _, n in ok))

    def test_ventas_emisor_incorrecto(self):
        c = clave(7, 9, ruc="1792222222001")
        md = md_lote(factura_md("v.pdf", c, "001-001-000000009", "07/08/2026", 10, 1.5, 11.5, ruc="1792222222001"))
        ruta = Path(self.tmp.name) / "v.md"
        ruta.write_text(md, encoding="utf-8")
        p = Parametros("venta", RUC, "EMPRENDEDORA PRUEBA", DESDE, HASTA)
        regs = validar_lote([extraer_campos(t, n) for n, t in leer_md(ruta).items()], p)
        self.assertEqual(regs[0]["estado"], "INVÁLIDA ❌")
        self.assertIn("RUC EMISOR INCORRECTO", " ".join(regs[0]["observaciones"]))
        wb = construir(regs, p, "PRUEBA", "Ventas_Ago2026_PRUEBA.xlsx", "Agosto 2026", "PRUE", 8)
        self.assertEqual(wb.sheetnames[0], "Registro de Ventas")
        self.assertEqual(wb["Registro de Ventas"]["A1"].value, "ANEXO MENSUAL DE INGRESOS")

    def test_ventas_familiar_y_canal(self):
        c1, c2 = clave(7, 11, ruc=RUC), clave(8, 12, ruc=RUC)
        md = md_lote(
            factura_md("v1.pdf", c1, "001-001-000000011", "07/08/2026", 10, 0, 10, receptor="PEREZ LOPEZ JUAN", idr="1790123456", ruc=RUC),
            factura_md("v2.pdf", c2, "001-001-000000012", "08/08/2026", 20, 0, 20, receptor="LOPEZ MARTINEZ ANA", idr="1790123457", ruc=RUC))
        ruta = Path(self.tmp.name) / "v2.md"
        ruta.write_text(md, encoding="utf-8")
        p = Parametros("venta", RUC, "LOPEZ ANDRADE LUCIA", DESDE, HASTA)
        regs = validar_lote([extraer_campos(t, n) for n, t in leer_md(ruta).items()], p)
        self.assertEqual([r["estado"] for r in regs], ["VÁLIDA ✅", "VÁLIDA ✅"])
        self.assertEqual([r["numero"] for r in regs], ["001-001-000000011", "001-001-000000012"])  # orden secuencial
        self.assertIn("LOPEZ", " ".join(regs[0]["observaciones"]))
        wb = construir(regs, p, "PRUEBA", "V.xlsx", "Agosto 2026", "PRUE", 7, "PRODUCTOS", clientes_b2b=["MARTINEZ"])
        ws = wb["Registro de Ventas"]
        canales = [c.value for c in ws["I"] if c.value in ("VENTAS B2B", "VENTAS B2C")]
        self.assertEqual(canales, ["VENTAS B2C", "VENTAS B2B"])


def pdf_minimo(n_paginas):
    """PDF válido de n páginas (para probar pdfseparate/pdfunite sin dependencias)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            "<< /Type /Pages /Kids [" + " ".join(f"{3 + 2 * i} 0 R" for i in range(n_paginas)) + f"] /Count {n_paginas} >>"]
    for i in range(n_paginas):
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents {4 + 2 * i} 0 R >>")
        cont = f"BT /F1 12 Tf 20 100 Td (pagina {i + 1}) Tj ET"
        objs.append(f"<< /Length {len(cont)} >>\nstream\n{cont}\nendstream")
    out, offs = b"%PDF-1.4\n", []
    for k, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{k} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode()
    return out


class TestPdfVariasFacturas(unittest.TestCase):
    def test_dividir_paginas(self):
        c1, c2 = clave(3, 1), clave(4, 2)
        paginas = [f"FACTURA\n{c1}\nitems", "", f"FACTURA\n{c2}\nitems", "VALOR TOTAL 10.00", "", ""]
        grupos = dividir_paginas(paginas, "todo.pdf")
        self.assertEqual([g[0] for g in grupos], ["todo.pdf#p01", "todo.pdf#p03-04"])
        self.assertIn("VALOR TOTAL 10.00", grupos[1][1])      # la página de totales se anexa a su factura
        self.assertNotIn(c2, grupos[0][1])

    @unittest.skipUnless(shutil.which("pdfseparate") and shutil.which("pdfunite") and shutil.which("pdfinfo"), "requiere poppler")
    def test_extraer_paginas(self):
        pdf = pdf_minimo(4)
        for desde, hasta in ((2, 2), (2, 3)):
            sub = extraer_paginas(pdf, desde, hasta)
            with tempfile.NamedTemporaryFile(suffix=".pdf") as f:
                f.write(sub)
                f.flush()
                info = subprocess.run(["pdfinfo", f.name], capture_output=True, text=True).stdout
            self.assertIn(f"Pages:           {hasta - desde + 1}", info)

    @unittest.skipUnless(shutil.which("pdfseparate") and shutil.which("pdfunite"), "requiere poppler")
    def test_renombrar_desde_pdf_unico(self):
        regs = [{"n": 1, "archivo": "todo.pdf#p01", "numero": "001-001-000000001", "receptor": "CLIENTE UNO", "modo": "venta"},
                {"n": 2, "archivo": "todo.pdf#p03-04", "numero": "001-001-000000002", "receptor": "CLIENTE DOS", "modo": "venta"}]
        with tempfile.TemporaryDirectory() as d:
            origen = Path(d) / "todo.pdf"
            origen.write_bytes(pdf_minimo(4))
            ok, falta = renombrar(regs, origen, Path(d) / "s.zip")
            self.assertEqual(falta, [])
            with zipfile.ZipFile(Path(d) / "s.zip") as z:
                self.assertEqual(sorted(z.namelist()), ["01_CLIENTE UNO_001-001-000000001.pdf", "02_CLIENTE DOS_001-001-000000002.pdf"])


class TestReglasVarias(unittest.TestCase):
    def test_huecos_numeracion(self):
        regs = [{"numero": f"002-100-{n:09d}"} for n in (95, 96, 97, 98, 100, 101)] + [{"numero": "001-001-000000005"}]
        self.assertEqual(huecos_numeracion(regs), {"002-100": ["002-100-000000099"]})
        self.assertEqual(huecos_numeracion([{"numero": "002-100-000000001"}, {"numero": "002-100-000000002"}]), {})

    def test_clasificacion_gasto(self):
        self.assertEqual(clasificar_gasto("A2 STAND - VENDEDOR 1M"), "GASTO VENTAS")
        self.assertEqual(clasificar_gasto("Servicio de suministro de dominio .ec para pagina web"), "GASTO VENTAS")
        self.assertEqual(clasificar_gasto("HONORARIOS PROFESIONALES"), "GASTO ADMINISTRATIVO")
        self.assertEqual(clasificar_gasto("FUNDA PLANA 15X20"), "COSTO OPERATIVO")


class TestConciliar(unittest.TestCase):
    def test_conciliar_detecta_faltas_y_errores(self):
        from fonquito.conciliar import conciliar
        regs = [{"n": 1, "numero": "001-001-000000001", "fecha_date": date(2026, 9, 1), "subtotal": 10.0, "iva": 0.0, "total": 10.0,
                 "id_receptor": "1790000000001", "receptor": "A", "estado": "VÁLIDA ✅", "duplicado_de": None},
                {"n": 2, "numero": "001-001-000000002", "fecha_date": date(2026, 9, 2), "subtotal": 20.0, "iva": 0.0, "total": 20.0,
                 "id_receptor": "1790000000002", "receptor": "B", "estado": "VÁLIDA ✅", "duplicado_de": None},
                {"n": 3, "numero": "001-001-000000003", "fecha_date": date(2026, 9, 3), "subtotal": 5.0, "iva": 0.0, "total": 5.0,
                 "id_receptor": "1790000000003", "receptor": "C", "estado": "VÁLIDA ✅", "duplicado_de": None}]
        filas = [{"# FACTURA": "001-001-000000001", "FECHA": "01/09/2026", "RUC": "1790000000001", "SUBTOTAL": 10, "IVA": 0, "TOTAL": 10},
                 {"# FACTURA": "001-001-000000002", "FECHA": "02/09/2026", "RUC": "1790000000002", "SUBTOTAL": 20.5, "IVA": 0, "TOTAL": 20},
                 {"# FACTURA": "001-001-000000009", "FECHA": "09/09/2026", "RUC": "1", "SUBTOTAL": 1, "IVA": 0, "TOTAL": 1}]
        res = conciliar(regs, filas, "venta")
        niveles = [(n, num) for n, _, num, _ in res]
        self.assertIn(("ERROR", "001-001-000000002"), niveles)     # subtotal distinto
        self.assertIn(("FALTA", "001-001-000000003"), niveles)     # no digitada
        self.assertIn(("SOBRA", "001-001-000000009"), niveles)     # sin factura
        self.assertNotIn(("ERROR", "001-001-000000001"), niveles)  # coincide


if __name__ == "__main__":
    unittest.main()
