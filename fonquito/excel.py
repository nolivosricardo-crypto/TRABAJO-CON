"""Genera el Excel institucional de validación (Registro + Validación Módulo 11 + Renombrado)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .renombrar import nombre_nuevo

HDR = PatternFill("solid", fgColor="2E75B6")
BAND = PatternFill("solid", fgColor="DDEBF7")
TOT = PatternFill("solid", fgColor="BDD7EE")
YEL = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
_t = Side(style="thin", color="B7C6D9")
BRD = Border(left=_t, right=_t, top=_t, bottom=_t)


def _f(b=False, sz=10, c="000000"):
    return Font(name="Arial", bold=b, size=sz, color=c)


def _fill(estado, i):
    if estado.startswith("INV"):
        return RED
    if estado.startswith("REV"):
        return YEL
    return BAND if i % 2 == 0 else None


def _yn(b):
    return "N/A" if b is None else ("✅" if b else "❌")


def _hdr(ws, fila, titulos):
    for j, h in enumerate(titulos, 1):
        c = ws.cell(fila, j, h)
        c.font, c.fill, c.border = _f(True, 10, "FFFFFF"), HDR, BRD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def construir(regs, params, emprendimiento, archivo_salida, mes_etiqueta, abrev=None, mes_contrato=None,
              categoria_venta=None, notas_extra=None, clientes_b2b=None, productos=None, conciliacion=None, ingreso_proyecto="APLICA"):
    """Crea el libro. `params` es validar.Parametros. Devuelve el Workbook."""
    gasto = params.tipo == "gasto"
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro de Gastos" if gasto else "Registro de Ventas"
    ws["A1"] = "REGISTRO DE GASTOS" if gasto else "ANEXO MENSUAL DE INGRESOS"
    ws["A1"].font = _f(True, 14)
    meta = [("Emprendimiento" if gasto else "NOMBRE DEL PROYECTO", emprendimiento),
            ("Código del fondo" if gasto else "CODIGO DEL PROYECTO", f"FQ2025-{emprendimiento}"),
            ("Emprendedora" if gasto else "RAZON SOCIAL", params.emprendedora),
            ("RUC/CI del beneficiario (receptor)" if gasto else "RUC emisor", params.ruc),
            ("Período válido", f"{params.desde:%d/%m/%Y} al {params.hasta:%d/%m/%Y}"),
            ("Archivo", archivo_salida)]
    if not gasto and abrev and mes_contrato:
        meta.append(("No", f"ING-EXPOST-{abrev}-{int(mes_contrato):02d}"))
    for k, (a, b) in enumerate(meta, 2):
        ws.cell(k, 1, a).font = _f(True)
        ws.cell(k, 3, b).font = _f()
    HR = len(meta) + 3
    cols = (["No.", "FECHA", "RUC/DNI", "# FACTURA", "PROVEEDOR", "PRODUCTO O SERVICIO", "ADQUISICION", "TIPO GASTO",
             "SUBTOTAL", "IVA", "TOTAL", "ESTADO"] if gasto else
            ["No.", "FECHA", "RUC", "# FACTURA", "CLIENTE", "PRODUCTO O SERVICIO", "ORIGEN", "INGRESO PROYECTO",
             "CANAL", "SUBTOTAL", "IVA", "TOTAL", "ESTADO"])
    _hdr(ws, HR, cols)
    num_cols = (9, 10, 11) if gasto else (10, 11, 12)
    est_col = len(cols)
    r0 = HR + 1
    for i, r in enumerate(regs):
        R = r0 + i
        if gasto:
            vals = [r["n"], r["fecha_date"], r.get("ruc_clave") or r.get("ruc_emisor"), r.get("numero"),
                    r.get("proveedor"), r.get("descripcion"), "NACIONAL", r["tipo_gasto"],
                    r["subtotal"], r["iva"], r["total"], r["estado"]]
        else:
            b2b = any(k.lower() in (r.get("receptor") or "").lower() or k in (r.get("id_receptor") or "")
                      for k in (clientes_b2b or []))
            etiquetas = [v for k, v in (productos or {}).items() if k.upper() in (r.get("descripcion") or "").upper()]
            producto = " / ".join(dict.fromkeys(etiquetas)) or categoria_venta or r.get("descripcion")
            vals = [r["n"], r["fecha_date"], r.get("id_receptor"), r.get("numero"), r.get("receptor"),
                    producto, "NACIONAL", ingreso_proyecto,
                    "VENTAS B2B" if b2b else "VENTAS B2C",
                    r["subtotal"], r["iva"], r["total"], r["estado"]]
        fill = _fill(r["estado"], i)
        if not gasto and not r["estado"].startswith(("INV", "REV")) and not b2b:
            fill = YEL  # canal asumido (B2C) a confirmar
        for j, v in enumerate(vals, 1):
            c = ws.cell(R, j, v)
            c.font, c.border = _f(), BRD
            c.alignment = Alignment(vertical="center", wrap_text=(j in (5, 6)))
            if fill:
                c.fill = fill
            if j == 2:
                c.number_format = "dd/mm/yyyy"
                c.alignment = Alignment(horizontal="center", vertical="center")
            if j in num_cols:
                c.number_format = "#,##0.00"
    last = r0 + len(regs) - 1
    TR = last + 1
    est = f"${chr(64 + est_col)}${r0}:${chr(64 + est_col)}${last}"
    for j in range(1, est_col + 1):
        ws.cell(TR, j).fill, ws.cell(TR, j).border, ws.cell(TR, j).font = TOT, BRD, _f(True)
    ws.cell(TR, 1, "TOTAL GASTOS (excluye duplicados e inválidas)" if gasto else "TOTAL FACTURACION INGRESOS MES").font = _f(True)
    letras = {c: chr(64 + c) for c in num_cols}
    for c in num_cols:
        L = letras[c]
        ws.cell(TR, c, f'=SUMIF({est},"<>INVÁLIDA ❌",{L}{r0}:{L}{last})').number_format = "#,##0.00"
        ws.cell(TR, c).font = _f(True)
        ws.cell(TR + 1, c, f'=SUMIF({est},"INVÁLIDA ❌",{L}{r0}:{L}{last})').number_format = "#,##0.00"
        ws.cell(TR + 2, c, f'=SUMIF({est},"REVISAR ⚠️",{L}{r0}:{L}{last})').number_format = "#,##0.00"
    ws.cell(TR + 1, 1, "Duplicados / inválidas no contabilizadas").font = _f()
    ws.cell(TR + 2, 1, "Marcadas REVISAR (incluidas en el total)").font = _f()
    notas = ["Filas amarillas = REVISAR (gastos) o canal asumido (ventas); filas rojas = inválida o duplicado, no se contabiliza.",
             "Diferencias de $0,01 por redondeo del emisor se consideran correctas.",
             ("Los comprobantes con C.I. (10 dígitos) en vez de RUC coinciden con el RUC del beneficiario." if gasto else
              "Las ventas a clientes con apellido en común con la emprendedora son legítimas; se dejan anotadas en la hoja 'Validación Módulo 11'."),
             "PRODUCTO O SERVICIO y TIPO GASTO se infieren de la descripción de la factura; confirmar con la emprendedora."
             if gasto else "CANAL: B2B para los clientes indicados por el técnico; el resto se asume VENTAS B2C (filas amarillas) hasta confirmar con la emprendedora.",
             "Detalle de cada validación en la hoja 'Validación Módulo 11'."] + (notas_extra or [])
    ws.cell(TR + 4, 1, "NOTAS").font = _f(True)
    for k, t in enumerate(notas, 1):
        ws.cell(TR + 4 + k, 1, f"{k}. {t}").font = _f(sz=9)
    for col, w in zip("ABCDEFGHIJKLM", [6, 12, 16, 20, 38, 48, 14, 22, 14, 12, 10, 12, 14]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[HR].height = 30
    ws.freeze_panes = ws.cell(r0, 1)

    # ---- Hoja 2
    w2 = wb.create_sheet("Validación Módulo 11")
    w2["A1"] = f"VALIDACIÓN MÓDULO 11 – {'GASTOS' if gasto else 'VENTAS'} {emprendimiento} ({mes_etiqueta})"
    w2["A1"].font = _f(True, 14)
    t2 = ["N°", "Clave de acceso", "Fecha extraída", "RUC extraído", "Factores Módulo 11", "Dígito calculado",
          "Dígito declarado", "Módulo 11 ✅/❌", "Fecha ✅/❌", "RUC ✅/❌"]
    if gasto:
        t2.append("Receptor ✅/❌")
    t2 += ["Monto ✅/❌", "Observación"]
    HR2 = 3
    _hdr(w2, HR2, t2)
    for i, r in enumerate(regs):
        R = HR2 + 1 + i
        suma, res = r["suma"], r["residuo"]
        factores = (f"Pesos 2-3-4-5-6-7 (derecha a izquierda); suma = {suma}; residuo = {res}; dígito = "
                    f"{0 if res == 0 else 1 if res == 1 else f'11 − {res} = {11 - res}'}") if suma is not None else "N/A"
        c = r["checks"]
        vals = [r["n"], r.get("clave") or "N/A", r["fecha_clave"], r["ruc_clave"], factores, r["dv_calc"], r["dv_decl"],
                _yn(c["m11"]), _yn(c["fecha"]), _yn(c["ruc"])]
        if gasto:
            vals.append(_yn(c["receptor"]))
        vals += [_yn(c["monto"]), "; ".join(r["observaciones"])]
        fill = _fill(r["estado"], i)
        for j, v in enumerate(vals, 1):
            cell = w2.cell(R, j, v)
            cell.font, cell.border = _f(sz=9 if j in (2, 5, len(vals)) else 10), BRD
            cell.alignment = Alignment(vertical="center", wrap_text=(j in (5, len(vals))),
                                       horizontal=("center" if 6 <= j < len(vals) or j in (1, 3) else None))
            if fill:
                cell.fill = fill
            if j == 3:
                cell.number_format = "dd/mm/yyyy"
            if j in (2, 4):
                cell.number_format = "@"
    l2 = HR2 + len(regs)
    est1 = f"'{ws.title}'!{est}"
    S = l2 + 3
    w2.cell(S - 1, 1, "CUADRO RESUMEN").font = _f(True, 12)
    filas = [("Total comprobantes procesados", f"=COUNTA({est1})"),
             ("Válidos (VÁLIDA ✅)", f'=COUNTIF({est1},"VÁLIDA ✅")'),
             ("Con incoherencias o a revisar (REVISAR ⚠️ + INVÁLIDA ❌)", f'=COUNTIF({est1},"REVISAR ⚠️")+COUNTIF({est1},"INVÁLIDA ❌")'),
             ("   de los cuales: REVISAR ⚠️", f'=COUNTIF({est1},"REVISAR ⚠️")'),
             ("   de los cuales: INVÁLIDA ❌ (incluye duplicados)", f'=COUNTIF({est1},"INVÁLIDA ❌")'),
             ("Monto total ejecutado (total con IVA, sin inválidas)" if gasto else "Monto total facturado (con IVA)",
              f"='{ws.title}'!{letras[num_cols[2]]}{TR}"),
             ("Ilegibles o incompletos", sum(1 for r in regs if r["advertencias"]))]
    for k, (a, b) in enumerate(filas):
        w2.cell(S + k, 1, a).font = _f(True)
        cb = w2.cell(S + k, 6, b)
        cb.font, cb.fill = _f(True), TOT
        for j in range(1, 7):
            w2.cell(S + k, j).border = BRD
        if k == 5:
            cb.number_format = '"$"#,##0.00'
        w2.merge_cells(start_row=S + k, start_column=1, end_row=S + k, end_column=5)
    d0 = S + len(filas) + 1
    w2.cell(d0, 1, "Incoherencias y motivos").font = _f(True, 11)
    k = 0
    for r in regs:
        if r["estado"].startswith("VÁL"):
            continue
        k += 1
        w2.cell(d0 + k, 1, f"N° {r['n']} ({r['archivo']}) – {r['estado']}: {'; '.join(r['observaciones'])}").font = _f(sz=9)
    anchos = [6, 52, 12, 16, 52, 10, 10, 12, 10, 10] + ([12] if gasto else []) + [10, 70]
    for idx, w in enumerate(anchos, 1):
        w2.column_dimensions[chr(64 + idx)].width = w
    w2.row_dimensions[HR2].height = 32
    w2.freeze_panes = w2.cell(HR2 + 1, 3)

    # ---- Conciliación con el formato institucional (opcional)
    if conciliacion is not None:
        w4 = wb.create_sheet("Conciliación formato")
        w4["A1"] = "CONCILIACIÓN: FORMATO INSTITUCIONAL vs FACTURAS"
        w4["A1"].font = _f(True, 14)
        _hdr(w4, 3, ["Nivel", "N°", "# Factura", "Detalle"])
        colores = {"ERROR": RED, "FALTA": RED, "SOBRA": YEL, "INFO": None}
        for i, (nivel, n, num, msg) in enumerate(conciliacion):
            for j, v in enumerate([nivel, n, num, msg], 1):
                c = w4.cell(4 + i, j, v)
                c.font, c.border = _f(), BRD
                c.alignment = Alignment(vertical="center", wrap_text=(j == 4))
                if colores.get(nivel):
                    c.fill = colores[nivel]
        if not any(x[0] != "INFO" for x in conciliacion):
            w4.cell(5 + len(conciliacion), 1, "Sin diferencias: el formato coincide con las facturas.").font = _f(True)
        for col, w in zip("ABCD", [10, 6, 22, 110]):
            w4.column_dimensions[col].width = w

    # ---- Hoja 3
    w3 = wb.create_sheet("Renombrado")
    w3["A1"] = "RENOMBRADO DE FACTURAS (<No.>_<proveedor/cliente>_<n° factura>.pdf)"
    w3["A1"].font = _f(True, 12)
    _hdr(w3, 3, ["No.", "Archivo original", "Nombre nuevo", "Estado"])
    for i, r in enumerate(regs):
        fill = _fill(r["estado"], i)
        for j, v in enumerate([r["n"], r["archivo"], nombre_nuevo(r), r["estado"]], 1):
            c = w3.cell(4 + i, j, v)
            c.font, c.border = _f(), BRD
            if fill:
                c.fill = fill
    for col, w in zip("ABCD", [6, 30, 74, 14]):
        w3.column_dimensions[col].width = w
    return wb
